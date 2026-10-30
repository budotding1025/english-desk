# -*- coding: utf-8 -*-
"""Build U3 school practice Word/PDF from screenshots (exam-from-photo)."""
from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

ROOT = Path(__file__).resolve().parent
CROPS = ROOT / "source" / "crops" / "u3"
OUT = ROOT


def font(run, name="宋体", size=11, bold=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold


def P(doc, text, *, size=11, bold=False, center=False, before=0, after=4, indent=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    if indent is not None:
        p.paragraph_format.first_line_indent = Cm(indent)
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    font(r, size=size, bold=bold)
    return p


def S(doc, text):
    P(doc, text, size=12, bold=True, before=10, after=6)


def footer(doc, page, total=4):
    P(doc, f"英语 · 四上 · 第三单元  第 {page} 页（共 {total} 页）", size=9, center=True, before=12, after=2)


def add_pic(doc, path: Path | None, width_cm=15.5):
    if not path or not path.exists():
        P(doc, "【图缺失】", size=10, center=True)
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    p.add_run().add_picture(str(path), width=Cm(width_cm))


def new_doc():
    doc = Document()
    for s in doc.sections:
        s.top_margin = Cm(1.6)
        s.bottom_margin = Cm(1.4)
        s.left_margin = Cm(1.8)
        s.right_margin = Cm(1.8)
    return doc


def build_paper() -> Path:
    doc = new_doc()
    P(doc, "英语 · 四上 · 第三单元练习", size=16, bold=True, center=True, after=8)
    P(doc, "小学________ 四年级____班  姓名____________  成绩________", size=11, center=True, after=10)
    P(doc, "听力部分", size=13, bold=True, center=True, after=8)

    S(doc, "一、听句子，选择正确的图片，在方框里画“√”。")
    add_pic(doc, CROPS / "sec1_full.jpg", 15.5)
    P(doc, "（上图为原卷附图；共 6 小题，在对应方框内画 √）", size=9, center=True, after=6)

    S(doc, "二、听对话，根据听到的顺序，用数字 1—5 给下列图片排序。")
    add_pic(doc, CROPS / "sec2_full.jpg", 15.5)

    S(doc, "三、听句子，选择正确的答语，将序号写在题前括号内。")
    reply3 = [
        [("A", "You're welcome."), ("B", "It's OK.")],
        [("A", "No, I don't."), ("B", "Sure, no problem.")],
        [("A", "Don't worry. You can use mine."), ("B", "I can fix it.")],
        [("A", "Not at all."), ("B", "It's OK.")],
        [("A", "Come in, please."), ("B", "No, you can't.")],
    ]
    for i, opts in enumerate(reply3, 1):
        P(doc, f"(  ) {i}.", size=11, bold=True, after=1)
        for lab, t in opts:
            P(doc, f"      {lab}. {t}", size=10, after=1)

    footer(doc, 1)
    doc.add_page_break()

    S(doc, "四、Henry 和同学们在班里做值日，他们是怎么分工的呢？听短文，将人物与其所负责的内容连线。")
    add_pic(doc, CROPS / "sec4_full.jpg", 15.5)
    P(doc, "人物：Mary / Jane / Henry / Jack / Lily", size=10, after=8)

    S(doc, "五、听对话，选择正确答案，将序号写在题前括号内。")
    q5 = [
        (
            "Where does this conversation happen?",
            [("A", "On a bus."), ("B", "In a taxi."), ("C", "On a subway.")],
        ),
        (
            "What is the next stop?",
            [
                ("A", "Beijing Railway Station."),
                ("B", "Beijing Airport."),
                ("C", "Dongzhimen Station."),
            ],
        ),
        (
            "Where does the woman want to go?",
            [
                ("A", "Beijing Railway Station."),
                ("B", "Dongzhimen Station."),
                ("C", "Not mentioned."),
            ],
        ),
        (
            "Which line should the woman take at Dongzhimen Station?",
            [("A", "Line 1."), ("B", "Line 2."), ("C", "Line 8.")],
        ),
    ]
    for i, (q, opts) in enumerate(q5, 1):
        P(doc, f"(  ) {i}. {q}", size=11, bold=True, after=2)
        for lab, t in opts:
            P(doc, f"      {lab}. {t}", size=10, after=1)

    P(doc, "笔试部分", size=13, bold=True, center=True, before=12, after=8)

    S(doc, "六、读单词，选出不同类的一项，将序号写在题前括号内。")
    odd = [
        "A. garden   B. push   C. stand   D. open",
        "A. polite   B. nice   C. line   D. clean",
        "A. problem   B. way   C. floor   D. find",
        "A. pass   B. use   C. bus   D. borrow",
        "A. excuse   B. hungry   C. sweep   D. wait",
    ]
    for i, line in enumerate(odd, 1):
        P(doc, f"(  ) {i}. {line}", size=11, after=4)

    footer(doc, 2)
    doc.add_page_break()

    S(doc, "七、读句子，选择与句子描述相符的词，将序号写在题前括号内。")
    for line in [
        "A. sweep",
        "B. favour",
        "C. line",
        "D. a minute",
        "E. problem",
    ]:
        P(doc, line, size=10, after=1)
    qs7 = [
        "Do something good to help others.",
        "Sixty seconds.",
        "Something is hard to do.",
        "To clean the dust from the floor.",
        "A row of people.",
    ]
    for i, q in enumerate(qs7, 1):
        P(doc, f"(  ) {i}. {q}", size=11, after=4)

    S(doc, "八、读对话，从方框中选择合适的句子补全对话，将序号写在横线上。")
    for line in [
        "A. Can I also get a banana cake?",
        "B. What can I do for you?",
        "C. Here you are.",
        "D. You are welcome.",
        "E. Here's the money. Thank you.",
    ]:
        P(doc, line, size=10, after=1)
    for line in [
        "Baker: Hi, 1. ________",
        "Amy: Excuse me, can I have two chocolate cakes?",
        "Baker: Sure, anything else?",
        "Amy: 2. ________",
        "Baker: Certainly. 3. ________",
        "Amy: How much are they?",
        "Baker: They're 30 yuan.",
        "Amy: 4. ________",
        "Baker: 5. ________ Would you like a bag?",
        "Amy: Yes, please.",
    ]:
        P(doc, line, size=11, after=3)

    S(doc, "九、看图读句子，选择与情景相符的句子，将序号写在题前括号内。")
    add_pic(doc, CROPS / "sec9_full.jpg", 15.5)
    q9 = [
        [
            ("A", "I'm sorry."),
            ("B", "Excuse me. Could you say it again?"),
            ("C", "How are you?"),
        ],
        [
            ("A", "May I borrow your ping-pong bat?"),
            ("B", "May I borrow your book?"),
            ("C", "May I come in?"),
        ],
        [
            ("A", "OK."),
            ("B", "See you."),
            ("C", "Not at all."),
        ],
        [
            ("A", "Could you please pass me the juice?"),
            ("B", "Have some juice, please."),
            ("C", "Give me some water."),
        ],
    ]
    for i, opts in enumerate(q9, 1):
        P(doc, f"(  ) {i}.", size=11, bold=True, after=1)
        for lab, t in opts:
            P(doc, f"      {lab}. {t}", size=10, after=1)

    footer(doc, 3)
    doc.add_page_break()

    S(doc, "十、以下是 George Middle School 的规则。阅读短文，完成任务。")
    P(doc, "The following rules are from George Middle School.", size=11, after=4)
    rules = [
        "Please wear school uniforms every day at school. And don't be late for class.",
        "Please walk on the right in the hallways (走廊). And don't run with your friends.",
        "Please keep the desks clean in the classroom. And don't chat or sleep in class.",
        "Please keep the books clean in the school library. And don't take them out.",
        "No eating or drinking in the computer lab. Turn off the computers before you leave.",
    ]
    for r in rules:
        P(doc, r, size=11, after=3)

    P(doc, "任务一：根据短文内容，判断句子正误，正确的写“T”，错误的写“F”。", size=11, bold=True, before=8, after=6)
    for i, s in enumerate(
        [
            "There are four places mentioned (提到) in the rules.",
            "The students can wear whatever they like at school.",
            "The students can take the book out from the school library.",
            "The students cannot eat or drink in the computer lab.",
        ],
        1,
    ):
        P(doc, f"(  ) {i}. {s}", size=11, after=4)

    P(doc, "任务二：根据短文内容，填写表格。", size=11, bold=True, before=8, after=6)
    table = doc.add_table(rows=5, cols=3)
    headers = ["Place", "Do", "Don't"]
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        for p in cell.paragraphs:
            for run in p.runs:
                font(run, size=10, bold=True)
    rows = [
        ("In the hallways", "walk on the right", "don't run with friends"),
        ("In the classroom", "keep the desks clean", "don't chat or sleep"),
        ("In the school library", "keep the books clean", "don't take them out"),
        ("In the computer lab", "turn off the computers", "no eating or drinking"),
    ]
    # blank student table (answers go in answer key; paper leaves blanks for 1–7 style)
    blanks = [
        ("In the hallways", "1. ______________", "2. ______________"),
        ("In the classroom", "3. ______________", "4. ______________"),
        ("In the school library", "5. ______________", "6. ______________"),
        ("In the computer lab", "7. ______________", "No eating or drinking"),
    ]
    for ri, row in enumerate(blanks, 1):
        for ci, val in enumerate(row):
            cell = table.rows[ri].cells[ci]
            cell.text = val
            for p in cell.paragraphs:
                for run in p.runs:
                    font(run, size=10)

    footer(doc, 4)

    path = OUT / "U3_第三单元练习.docx"
    doc.save(path)
    shutil.copy(path, OUT / "U3_Unit3_Practice.docx")
    return path


def build_answers() -> Path:
    doc = new_doc()
    P(doc, "英语 · 四上 · 第三单元练习 · 参考答案", size=16, bold=True, center=True, after=6)
    P(doc, "含完整英语听力原文 + 全部题目答案（听力据《四上听力材料》第三单元）", size=10, center=True, after=12)

    S(doc, "听力部分 · 原文与答案")

    P(doc, "一、听选图", size=12, bold=True, after=4)
    P(doc, "【Listening script】", size=10, bold=True, after=2)
    for i, t in enumerate(
        [
            "Could you please open the door for me?",
            "We should stand in line and get on the bus one by one.",
            "You want to ask the way. What will you say at first?",
            "Could you please pass me that book?",
            "I'd like to sweep the floor.",
            "The boy can fix the model plane.",
        ],
        1,
    ):
        P(doc, f"{i}. {t}", size=10, after=2)
    P(
        doc,
        "【Answers】Tick √ the picture that matches each sentence (6 items). "
        "对照练习卷/扫描卷附图勾选。",
        size=10,
        after=8,
    )

    P(doc, "二、听对话排序 1—5", size=12, bold=True, after=4)
    P(doc, "【Listening script】（heard in this order）", size=10, bold=True, after=2)
    for i, t in enumerate(
        [
            "Could you get me some water please? — Sure, no problem.",
            "Don't push, Henry. It's not polite. — I see. Sorry, Uncle.",
            "Excuse me, could I have some rice, please? — Sure, here you are.",
            "Could you help me hold these books please? — Sure, no problem.",
            "Excuse me, I can't find my pencil case. May I borrow a pencil? — Sure, here is one.",
        ],
        1,
    ):
        P(doc, f"{i}. {t}", size=10, after=2)
    P(doc, "【Answers】Write 1–5 on the pictures in the order you hear. 对照附图填写。", size=10, after=8)

    P(doc, "三、听选答语", size=12, bold=True, after=4)
    P(doc, "【Listening script】", size=10, bold=True, after=2)
    for i, t in enumerate(
        [
            "You are so kind. Thank you.",
            "Could you help me to sweep the floor please?",
            "My eraser is missing.",
            "I'm sorry I'm late.",
            "May I come in?",
        ],
        1,
    ):
        P(doc, f"{i}. {t}", size=10, after=2)
    P(doc, "【Answers】1A  2B  3A  4B  5A", size=11, bold=True, after=2)
    P(
        doc,
        "1 Thank you → You're welcome.  2 help sweep → Sure, no problem.  "
        "3 eraser missing → Don't worry. You can use mine.  "
        "4 I'm sorry I'm late → It's OK.  5 May I come in? → Come in, please.",
        size=9,
        after=8,
    )

    P(doc, "四、值日连线", size=12, bold=True, after=4)
    P(doc, "【Listening script】", size=10, bold=True, after=2)
    P(
        doc,
        "I am Henry. It's Wednesday today. Our team is on duty. Mary will clean the blackboard. "
        "Jane will tidy up the desks. I will sweep the floor, and Jack would like to help me sweep together. "
        "Lily helps us get some water and close all the windows. Everyone is happy to make the classroom clean.",
        size=10,
        after=4,
    )
    P(doc, "【Answers】", size=10, bold=True, after=2)
    P(doc, "Mary → clean the blackboard", size=10, after=1)
    P(doc, "Jane → tidy up the desks", size=10, after=1)
    P(doc, "Henry → sweep the floor", size=10, after=1)
    P(doc, "Jack → help Henry sweep", size=10, after=1)
    P(doc, "Lily → get some water and close all the windows", size=10, after=8)

    P(doc, "五、听对话选择", size=12, bold=True, after=4)
    P(doc, "【Listening script】", size=10, bold=True, after=2)
    for line in [
        "Lady A: Excuse me, could you tell me if this train stops at Dongzhimen Station?",
        "Lady B: Yes, it does. The next stop is Dongzhimen Station.",
        "Lady A: Thank you! And I want to go to Beijing Railway Station. Would you mind showing me how to get there?",
        "Lady B: Sure. Let me see — At Dongzhimen Station, you should take Line 2 to Beijing Railway Station.",
        "Lady A: That's very helpful. Thank you very much.",
        "Lady B: You're welcome. Have a nice day!",
    ]:
        P(doc, line, size=9, after=1)
    P(doc, "【Answers】1C  2C  3A  4B", size=11, bold=True, after=2)
    P(
        doc,
        "1 on a subway/train；2 next stop Dongzhimen；3 Beijing Railway Station；4 Line 2。",
        size=9,
        after=10,
    )

    S(doc, "笔试部分")
    P(doc, "六、异类词：1A  2C  3D  4C  5B", size=11, after=4)
    P(
        doc,
        "1 garden 名词，其余动词；2 line 名词，其余形容/动词；3 find 动词，其余名词；"
        "4 bus 名词，其余动词；5 hungry 形容词，其余动词。",
        size=9,
        after=6,
    )

    P(doc, "七、词义匹配：1B  2D  3E  4A  5C", size=11, after=6)

    P(doc, "八、补全对话：1B  2A  3C  4E  5D", size=11, after=6)

    P(doc, "九、看图选句：1B  2A  3C  4A", size=11, after=4)
    P(
        doc,
        "参考：1 请再说一遍 → Excuse me. Could you say it again?；"
        "2 借球拍 → May I borrow your ping-pong bat?；"
        "3 回应感谢 → Not at all.；4 请递果汁 → Could you please pass me the juice?",
        size=9,
        after=6,
    )

    P(doc, "十·任务一：1T  2F  3F  4T", size=11, after=4)
    P(
        doc,
        "1 走廊/教室/图书馆/电脑室共四处；2 须穿校服；3 书不可带出；4 电脑室禁饮食。",
        size=9,
        after=6,
    )
    P(doc, "十·任务二（参考）：", size=11, bold=True, after=4)
    P(doc, "In the hallways：walk on the right / don't run with friends", size=10, after=2)
    P(doc, "In the classroom：keep the desks clean / don't chat or sleep", size=10, after=2)
    P(doc, "In the school library：keep the books clean / don't take them out", size=10, after=2)
    P(doc, "In the computer lab：turn off the computers / no eating or drinking", size=10, after=8)

    path = OUT / "U3_第三单元练习_答案.docx"
    doc.save(path)
    shutil.copy(path, OUT / "U3_Unit3_Practice_Answers.docx")
    return path


def export_pdfs(paths: list[Path]) -> None:
    import win32com.client

    word = win32com.client.Dispatch("Word.Application")
    word.Visible = False
    try:
        for docx in paths:
            pdf = docx.with_suffix(".pdf")
            d = word.Documents.Open(str(docx))
            d.SaveAs(str(pdf), FileFormat=17)
            d.Close(False)
            print("pdf", pdf.name)
    finally:
        word.Quit()


def build_scan_pdf() -> Path | None:
    try:
        from PIL import Image
    except ImportError:
        return None
    pages = []
    for name in ["u3-1a_hq", "u3-1b_hq", "u3-2a_hq", "u3-2b_hq"]:
        p = CROPS / f"{name}.jpg"
        if p.exists():
            pages.append(Image.open(p).convert("RGB"))
    if not pages:
        return None
    out = OUT / "U3_Unit3_Practice_Scan.pdf"
    pages[0].save(out, save_all=True, append_images=pages[1:], resolution=150)
    shutil.copy(out, OUT / "U3_第三单元练习_原卷扫描.pdf")
    print("scan pdf", out.name)
    return out


if __name__ == "__main__":
    paper = build_paper()
    ans = build_answers()
    print("docx", paper.name, ans.name)
    build_scan_pdf()
    export_pdfs(
        [
            paper,
            OUT / "U3_Unit3_Practice.docx",
            ans,
            OUT / "U3_Unit3_Practice_Answers.docx",
        ]
    )
    shutil.copy(OUT / "U3_Unit3_Practice.pdf", OUT / "U3_第三单元练习.pdf")
    shutil.copy(OUT / "U3_Unit3_Practice_Answers.pdf", OUT / "U3_第三单元练习_答案.pdf")
    print("done")
