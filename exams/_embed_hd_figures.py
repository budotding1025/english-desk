# -*- coding: utf-8 -*-
"""Rebuild U1/U2 practice Word/PDF with HD regenerated exam figures."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, Inches

ROOT = Path(__file__).resolve().parent
HD = ROOT / "papers" / "figures" / "hd"
OUT = ROOT

# Prefer jpg from GenerateImage assets copy
def fig(name: str) -> Path:
    for ext in (".jpg", ".png", ".jpeg"):
        p = HD / f"{name}{ext}"
        if p.exists():
            return p
    raise FileNotFoundError(name)


def set_run_font(run, name="宋体", size=11, bold=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold


def P(doc, text, *, size=11, bold=False, center=False, before=0, after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    set_run_font(r, size=size, bold=bold)
    return p


def S(doc, text):
    P(doc, text, size=12, bold=True, before=10, after=6)


def add_img(doc, path: Path, width_cm=3.2):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run()
    run.add_picture(str(path), width=Cm(width_cm))
    return p


def pair_row(doc, n, left: Path, right: Path, w=3.0):
    P(doc, f"{n}.", size=11, bold=True, after=2)
    table = doc.add_table(rows=2, cols=2)
    table.autofit = True
    # images
    for i, path in enumerate((left, right)):
        cell = table.rows[0].cells[i]
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = cell.paragraphs[0].add_run()
        run.add_picture(str(path), width=Cm(w))
    # labels
    for i, lab in enumerate(("□ A", "□ B")):
        cell = table.rows[1].cells[i]
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = cell.paragraphs[0].add_run(lab)
        set_run_font(r, size=10, bold=True)
    P(doc, "", size=6, after=4)


def underline_pair(doc, n, parts):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    r0 = p.add_run(f"(  ) {n}. ")
    set_run_font(r0, size=11)
    for text, under in parts:
        r = p.add_run(text)
        set_run_font(r, size=11)
        r.underline = under


def new_doc():
    doc = Document()
    for s in doc.sections:
        s.top_margin = Cm(1.6)
        s.bottom_margin = Cm(1.4)
        s.left_margin = Cm(1.8)
        s.right_margin = Cm(1.8)
    return doc


def footer(doc, unit, page, total=4):
    P(doc, f"英语 · 四上 · 第{unit}单元  第 {page} 页（共 {total} 页）· 高清图题版", size=9, center=True, before=10, after=2)


def build_u1():
    doc = new_doc()
    P(doc, "英语 · 四上 · 第一单元练习", size=16, bold=True, center=True, after=4)
    P(doc, "（图题已按原卷重绘高清版）", size=9, center=True, after=6)
    P(doc, "小学________ 四年级____班  姓名____________  成绩________", size=11, center=True, after=8)
    P(doc, "听力部分", size=13, bold=True, center=True, after=8)

    S(doc, "一、听句子，选择正确的图片，在方框里面画“√”。")
    for i in range(1, 7):
        pair_row(doc, i, fig(f"u1_L1_{i}A"), fig(f"u1_L1_{i}B"), w=3.0)

    S(doc, "二、听对话，根据听到的顺序，用数字 1–4 给下列图片排序。")
    table = doc.add_table(rows=2, cols=4)
    for i, lab in enumerate("ABCD"):
        cell = table.rows[0].cells[i]
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cell.paragraphs[0].add_run().add_picture(str(fig(f"u1_L2_{lab}")), width=Cm(2.8))
        c2 = table.rows[1].cells[i]
        c2.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = c2.paragraphs[0].add_run(f"{lab}. 排序：____")
        set_run_font(r, size=9)

    S(doc, "三、听句子，选择正确的答语，将序号写在题前括号内。")
    reply3 = [
        [("A", "I'm excited."), ("B", "I'm reading a book.")],
        [("A", "He can't find his cat."), ("B", "It is small and thin.")],
        [("A", "Let's go and play football."), ("B", "I'm angry because he broke my pencil.")],
        [("A", "I'm happy now."), ("B", "I don't like it.")],
        [("A", "She looks worried."), ("B", "She can't find her English book.")],
    ]
    for i, opts in enumerate(reply3, 1):
        P(doc, f"(  ) {i}.", bold=True, after=1)
        for lab, t in opts:
            P(doc, f"      {lab}. {t}", size=10, after=1)

    footer(doc, "一", 1)
    doc.add_page_break()

    S(doc, "四、听短文，判断句子正误，正确的画“√”，错误的画“×”。")
    for i, s in enumerate(
        [
            "Today is the first day of school.",
            "I am happy because I can play with friends and read interesting books at school.",
            "My mother is angry because I don't do homework.",
            "After school, I play games in the classroom.",
            "My mother becomes happy again after I clean my room.",
        ],
        1,
    ):
        P(doc, f"(  ) {i}. {s}", after=3)

    P(doc, "笔试部分", size=13, bold=True, center=True, before=10, after=8)
    S(doc, "五、判断下列单词划线部分发音是否相同，相同的画“√”，不相同的画“×”。")
    for i, parts in enumerate(
        [
            [("c", False), ("a", True), ("ke", False), ("  /  m", False), ("a", True), ("ke", False)],
            [("c", False), ("a", True), ("t", False), ("  /  ", False), ("a", True), ("ngry", False)],
            [("h", False), ("e", True), ("  /  sh", False), ("e", True)],
            [("b", False), ("e", True), ("d", False), ("  /  m", False), ("e", True)],
        ],
        1,
    ):
        underline_pair(doc, i, parts)

    S(doc, "六、从方框中选择与所给单词同类的单词，将序号写在横线上。")
    P(doc, "方框：A. sunny  B. worried  C. sad  D. fix  E. angry  F. cloudy  G. look  H. find  I. rainy", size=10, after=4)
    P(doc, "1. happy  ________  ________  ________", after=3)
    P(doc, "2. worry  ________  ________  ________", after=3)
    P(doc, "3. windy  ________  ________  ________", after=4)

    S(doc, "七、看图读句子，选择与图片相符的句子，将序号写在括号内。")
    for line in [
        "A. My aunt is angry, because the cat broke the vase (花瓶).",
        "B. The kids are excited, because they are going to the amusement park.",
        "C. The boy is sad. He can't find his red scarf.",
        "D. My mum is worried, because my grandma is ill.",
        "E. Peter is happy, because he plays basketball with his friends.",
    ]:
        P(doc, line, size=10, after=1)

    table = doc.add_table(rows=2, cols=5)
    for i in range(5):
        cell = table.rows[0].cells[i]
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cell.paragraphs[0].add_run().add_picture(str(fig(f"u1_W7_{i + 1}")), width=Cm(2.6))
        c2 = table.rows[1].cells[i]
        c2.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = c2.paragraphs[0].add_run(f"(  ) {i + 1}.")
        set_run_font(r, size=10)

    footer(doc, "一", 2)
    doc.add_page_break()

    S(doc, "八、读问句，从方框中选择正确的答语，将序号写在题前括号内。")
    for line in [
        "A. I'm happy.",
        "B. He can't find his dog.",
        "C. I'm better, Mum.",
        "D. Because Kevin broke his new model plane.",
        "E. He is brown with a black nose.",
    ]:
        P(doc, line, size=10, after=1)
    for i, q in enumerate(
        [
            "What's the matter with Maomao?",
            "How do you feel on the first school day?",
            "Why is Mike angry?",
            "What does Danny look like?",
            "How do you feel now?",
        ],
        1,
    ):
        P(doc, f"(  ) {i}. {q}", after=3)

    S(doc, "九、读对话，从方框中选择合适的句子补全对话，将序号写在横线上。")
    for line in [
        "A. Thank you so much.",
        "B. He will be angry.",
        "C. What's the matter?",
        "D. What does it look like?",
        "E. Look, it's just under your cap.",
    ]:
        P(doc, line, size=10, after=1)
    for line in [
        "A: Hi, John. You look sad. 1. ________",
        "B: I can't find my watch. My father bought it for me only yesterday. 2. ________",
        "A: Don't be sad. I can help you. Let's look for it together. 3. ________",
        "B: It's round in shape and black in colour.",
        "A: 4. ________",
        "B: Oh, that's it. 5. ________",
        "A: You are welcome.",
    ]:
        P(doc, line, after=2)

    S(doc, "十、阅读短文，完成任务。")
    P(doc, "I'm Happy", size=12, bold=True, center=True, after=4)
    passage = (
        "I'm happy when the sun is shining. I can go outside and play with my friends. "
        "We can ride our bikes, fly kites, or play catch(接球游戏). The warm sun makes me feel good. "
        "I'm happy when I'm with my family. We eat dinner together, talk and laugh. "
        "I'm happy when I read a good book. It's like having an adventure(冒险) in my mind. "
        "I'm happy when I help others and see a smile on his face."
    )
    P(doc, passage, after=6)
    footer(doc, "一", 3)
    doc.add_page_break()

    P(doc, "任务一：根据短文内容，选择正确答案。", bold=True, after=4)
    read_qs = [
        ("When am I happy?", [("A", "When it is raining."), ("B", "When the sun is shining."), ("C", "When it is snowing.")]),
        ("What can I do with friends when the sun is shining?", [("A", "Watch TV."), ("B", "Do homework."), ("C", "Ride bikes, fly kites or play catch.")]),
        ("How do I feel when I am with my family?", [("A", "Safe and loved."), ("B", "Angry."), ("C", "Sad.")]),
        ("What do I think of reading a good book?", [("A", "Boring."), ("B", "Difficult."), ("C", "Like having an adventure in the mind.")]),
        ("Why do I feel happy when helping others?", [("A", "Because I can get money."), ("B", "Because I can get new toys."), ("C", "Because I can see a smile on the friend's face.")]),
    ]
    for i, (q, opts) in enumerate(read_qs, 1):
        P(doc, f"(  ) {i}. {q}", bold=True, after=1)
        for lab, t in opts:
            P(doc, f"      {lab}. {t}", size=10, after=1)
    P(doc, "任务二：仿写 I'm ________ when _________________________________", bold=True, before=8, after=4)

    P(doc, "—— 参考答案（教师用）——", size=12, bold=True, center=True, before=14, after=6)
    for line in [
        "三：1A  2A  3B  4A  5B",
        "四：1√  2√  3×  4×  5√",
        "五：1√  2√  3√  4×",
        "六：1 B/C/E  2 D/G/H  3 A/F/I",
        "七：1B  2A  3E  4C  5D",
        "八：1B  2A  3D  4E  5C",
        "九：1C  2B  3D  4E  5A",
        "十·任务一：1B  2C  3A  4C  5C",
    ]:
        P(doc, line, size=9, after=2)
    footer(doc, "一", 4)

    path = OUT / "U1_Unit1_Practice.docx"
    doc.save(path)
    doc.save(OUT / "U1_第一单元练习.docx")
    return path


def build_u2():
    doc = new_doc()
    P(doc, "英语 · 四上 · 第二单元练习", size=16, bold=True, center=True, after=4)
    P(doc, "（图题已按原卷重绘高清版）", size=9, center=True, after=6)
    P(doc, "小学________ 四年级____班  姓名____________  成绩________", size=11, center=True, after=8)
    P(doc, "听力部分", size=13, bold=True, center=True, after=8)

    S(doc, "一、听句子，选择正确的图片，在方框里画“√”。")
    for i in range(1, 7):
        pair_row(doc, i, fig(f"u2_L1_{i}A"), fig(f"u2_L1_{i}B"), w=3.0)

    S(doc, "二、听对话，判断图片与所听内容是否相符，相符的画“√”，不相符的画“×”。")
    table = doc.add_table(rows=2, cols=5)
    for i in range(5):
        cell = table.rows[0].cells[i]
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cell.paragraphs[0].add_run().add_picture(str(fig(f"u2_L2_{i + 1}")), width=Cm(2.5))
        c2 = table.rows[1].cells[i]
        c2.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = c2.paragraphs[0].add_run(f"(  ) {i + 1}.")
        set_run_font(r, size=10)

    S(doc, "三、听问句，选择正确的答语，将序号写在题前括号内。")
    for i, opts in enumerate(
        [
            [("A", "I am Sara."), ("B", "Sorry, she is not home.")],
            [("A", "I can't remember the poem."), ("B", "She can't remember the poem.")],
            [("A", "She is doing homework."), ("B", "He is doing homework.")],
            [("A", "Sure."), ("B", "Sorry, I can't.")],
            [("A", "I'm drawing a picture."), ("B", "I'm going to the park tomorrow.")],
        ],
        1,
    ):
        P(doc, f"(  ) {i}.", bold=True, after=1)
        for lab, t in opts:
            P(doc, f"      {lab}. {t}", size=10, after=1)

    footer(doc, "二", 1)
    doc.add_page_break()

    S(doc, "四、听短文，判断句子正误，正确的写“T”，错误的写“F”。")
    for i, s in enumerate(
        [
            "Tina is my good friend.",
            "We go to school by car every day.",
            "Art is our favourite class.",
            "We help each other only in art class.",
            "We go to the zoo and fly kites on weekends.",
        ],
        1,
    ):
        P(doc, f"(  ) {i}. {s}", after=3)

    P(doc, "笔试部分", size=13, bold=True, center=True, before=10, after=8)
    S(doc, "五、判断下列单词画线部分发音是否相同，相同的画“√”，不相同的画“×”。")
    for i, parts in enumerate(
        [
            [("n", False), ("o", True), ("  /  n", False), ("o", True), ("t", False)],
            [("f", False), ("i", True), ("ve", False), ("  /  m", False), ("i", True), ("ce", False)],
            [("s", False), ("e", True), ("t", False), ("  /  s", False), ("i", True), ("t", False)],
            [("g", False), ("o", True), ("  /  ag", False), ("o", True)],
        ],
        1,
    ):
        underline_pair(doc, i, parts)

    S(doc, "六、读单词，选出不同类的一项，将序号写在题前括号内。")
    for i, line in enumerate(
        [
            "A. remember   B. help   C. hard   D. try",
            "A. sleeping   B. reading   C. telling   D. sing",
            "A. story   B. visit   C. book   D. gift",
            "A. good   B. call   C. hear   D. speak",
            "A. friend   B. mother   C. brother   D. grandpa",
        ],
        1,
    ):
        P(doc, f"(  ) {i}. {line}", after=3)

    S(doc, "七、看图读句子，选择与图片相符的句子，将序号写在括号内。")
    for line in [
        "A. The book is about airplane.",
        "B. Bill is sleeping. He is ill.",
        "C. You can draw a picture for a poem.",
        "D. Kevin has a bad cold.",
        "E. You can drink some warm water.",
    ]:
        P(doc, line, size=10, after=1)
    table = doc.add_table(rows=2, cols=5)
    for i in range(5):
        cell = table.rows[0].cells[i]
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cell.paragraphs[0].add_run().add_picture(str(fig(f"u2_W7_{i + 1}")), width=Cm(2.5))
        c2 = table.rows[1].cells[i]
        c2.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = c2.paragraphs[0].add_run(f"(  ) {i + 1}.")
        set_run_font(r, size=10)

    footer(doc, "二", 2)
    doc.add_page_break()

    S(doc, "八、读句子，从方框中选择正确的答语，将序号写在题前括号内。")
    for line in [
        "A. He is reading a science book.",
        "B. I can't remember the poem.",
        "C. Not so good.",
        "D. Sorry, she is sleeping.",
        "E. I often play football with him.",
        "F. Thank you so much.",
    ]:
        P(doc, line, size=10, after=1)
    for i, q in enumerate(
        [
            "What's the matter, Alice?",
            "May I speak to Lucy?",
            "What do you often do with him?",
            "How are you feeling?",
            "What is he reading?",
            "You should have a good rest.",
        ],
        1,
    ):
        P(doc, f"(  ) {i}. {q}", after=3)

    S(doc, "九、读对话，从方框中选择合适的句子补全对话，将序号写在横线上。")
    for line in [
        "A. I have some new story books.",
        "B. I'll tell her.",
        "C. Can you ask her to call me back?",
        "D. We can read books together.",
        "E. May I speak to Sophia?",
    ]:
        P(doc, line, size=10, after=1)
    for line in [
        "Emma: Hello, Auntie. This is Emma. 1. ________",
        "Mrs Smith: Sorry, Emma. She is not home.",
        "Emma: Oh, I see. 2. ________",
        "Mrs Smith: Sure. 3. ________",
        "Emma: Thank you, Auntie.",
        "Sophia: Hi, Emma! 4. ________ Do you want to read with me?",
        "Emma: Yes. Will you bring them to school tomorrow?",
        "Sophia: Certainly.",
        "Emma: Great! 5. ________",
    ]:
        P(doc, line, after=2)

    S(doc, "十、阅读短文，完成任务。")
    P(doc, "Friends Help Each Other", size=12, bold=True, center=True, after=4)
    P(
        doc,
        "I am Lucy. I have a great friend named Lily. We are in the same class. "
        "Last Monday, we had a big maths test. I was so nervous because I lost my ruler. "
        "Lily gave me her spare ruler. Lily is not good at drawing. I often help her draw. "
        "We always help each other and our friendship is getting stronger.",
        after=6,
    )
    footer(doc, "二", 3)
    doc.add_page_break()

    P(doc, "任务一：根据短文内容，选择正确答案。", bold=True, after=4)
    for i, (q, opts) in enumerate(
        [
            ("What is the name of the writer's friend?", [("A", "Lucy."), ("B", "Lily."), ("C", "Lisa.")]),
            ("What happened to the writer before the maths test?", [("A", "She lost the textbook."), ("B", "She lost the ruler."), ("C", "She lost the pencil.")]),
            ("What did Lily do when she saw the writer worried?", [("A", "She gave her another ruler."), ("B", "She helped write answers."), ("C", "She asked the teacher for help.")]),
            ("What is Lily not good at?", [("A", "Maths."), ("B", "Reading."), ("C", "Drawing.")]),
            ("How did the writer help Lily?", [("A", "She taught her maths."), ("B", "She taught her reading."), ("C", "She taught her how to draw.")]),
        ],
        1,
    ):
        P(doc, f"(  ) {i}. {q}", bold=True, after=1)
        for lab, t in opts:
            P(doc, f"      {lab}. {t}", size=10, after=1)
    P(doc, "任务二：写两句介绍自己的朋友。", bold=True, before=8, after=4)
    P(doc, "例：Lily is my good friend. I often teach her how to draw.", size=10, after=4)
    P(doc, "1. _______________________________________________", after=3)
    P(doc, "2. _______________________________________________", after=6)

    P(doc, "—— 参考答案（教师用）——", size=12, bold=True, center=True, before=10, after=6)
    for line in [
        "三：1B  2A  3B  4A  5A",
        "四：1T  2F  3T  4F  5F",
        "五：1×  2√  3×  4√",
        "六：1C  2D  3B  4A  5A",
        "七：1B  2A  3E  4C  5D",
        "八：1B  2D  3E  4C  5A  6F",
        "九：1E  2C  3B  4A  5D",
        "十·任务一：1B  2B  3A  4C  5C",
    ]:
        P(doc, line, size=9, after=2)
    footer(doc, "二", 4)

    path = OUT / "U2_Unit2_Practice.docx"
    doc.save(path)
    doc.save(OUT / "U2_第二单元练习.docx")
    return path


def export_pdfs(paths):
    import win32com.client

    word = win32com.client.Dispatch("Word.Application")
    word.Visible = False
    try:
        for docx in paths:
            pdf = docx.with_suffix(".pdf")
            d = word.Documents.Open(str(docx))
            d.SaveAs(str(pdf), FileFormat=17)
            d.Close(False)
            print("pdf", pdf.name)
    finally:
        word.Quit()


def write_fig_readme():
    text = """# 单元卷图题 · 高清重绘

来源：`Unit test/U1-1.png`、`U2-1.png` 原卷扫描裁切 + 提示词重绘。

## 目录

| 路径 | 说明 |
|------|------|
| `crops/` | 原卷裁切对照 |
| `hd/` | 高清新图（43 张，U1/U2 图题全覆盖） |
| `_crop_panels.py` | 裁切脚本 |

## 嵌入试卷

运行 `exams/_embed_hd_figures.py` 生成带图 Word/PDF：

- `exams/U1_Unit1_Practice.docx` / `.pdf`
- `exams/U2_Unit2_Practice.docx` / `.pdf`

图意已按 `exams/AUDIT.md` 核对（含红围巾、病床旁站立、U2 发抖感冒等）。
"""
    (ROOT / "papers" / "figures" / "README.md").write_text(text, encoding="utf-8")


if __name__ == "__main__":
    write_fig_readme()
    u1 = build_u1()
    u2 = build_u2()
    print("docx", u1.name, u2.name)
    export_pdfs([u1, u2])
    # refresh Chinese named pdfs
    import shutil

    shutil.copy(u1.with_suffix(".pdf"), OUT / "U1_第一单元练习.pdf")
    shutil.copy(u2.with_suffix(".pdf"), OUT / "U2_第二单元练习.pdf")
    print("done")
