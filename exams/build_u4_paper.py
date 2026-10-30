# -*- coding: utf-8 -*-
"""Build U4 school practice Word/PDF from OCR + photo crops (exam-from-photo)."""
from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

ROOT = Path(__file__).resolve().parent
CELLS = ROOT / "source" / "crops" / "cells"
HQ = ROOT / "source" / "crops"
OUT = ROOT


def font(run, name="宋体", size=11, bold=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold


def P(doc, text, *, size=11, bold=False, center=False, before=0, after=4, indent=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    if indent is not None:
        p.paragraph_format.first_line_indent = Cm(indent)
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    font(r, size=size, bold=bold)
    return p


def S(doc, text):
    P(doc, text, size=12, bold=True, before=10, after=6)


def footer(doc, page, total=4):
    P(doc, f"英语 · 四上 · 第四单元  第 {page} 页（共 {total} 页）", size=9, center=True, before=12, after=2)


def cell(name: str) -> Path | None:
    for ext in (".jpg", ".png", ".jpeg"):
        p = CELLS / f"{name}{ext}"
        if p.exists():
            return p
    return None


def add_pic(doc, path: Path | None, width_cm=3.0):
    if not path or not path.exists():
        P(doc, "【图缺失】", size=10, center=True)
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(str(path), width=Cm(width_cm))


def grid_pics(doc, names, labels=None, width=2.8, cols=4):
    labels = labels or [str(i + 1) for i in range(len(names))]
    row_n = (len(names) + cols - 1) // cols
    table = doc.add_table(rows=row_n * 2, cols=cols)
    table.autofit = True
    for i, name in enumerate(names):
        r, c = divmod(i, cols)
        img_cell = table.rows[r * 2].cells[c]
        lab_cell = table.rows[r * 2 + 1].cells[c]
        img_cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        path = cell(name)
        if path:
            img_cell.paragraphs[0].add_run().add_picture(str(path), width=Cm(width))
        lab_cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = lab_cell.paragraphs[0].add_run(f"{labels[i]}. (   )")
        font(run, size=10, bold=True)
    P(doc, "", size=4, after=2)


def new_doc():
    doc = Document()
    for s in doc.sections:
        s.top_margin = Cm(1.6)
        s.bottom_margin = Cm(1.4)
        s.left_margin = Cm(1.8)
        s.right_margin = Cm(1.8)
    return doc


def build_paper() -> Path:
    doc = new_doc()
    P(doc, "英语 · 四上 · 第四单元练习", size=16, bold=True, center=True, after=8)
    P(doc, "小学________ 四年级____班  姓名____________  成绩________", size=11, center=True, after=10)
    P(doc, "听力部分", size=13, bold=True, center=True, after=8)

    S(doc, "一、听句子，判断图片与所听内容是否相符，相符的写“T”，不相符的写“F”。")
    # prefer cleaned s1_* ; fallback s1r*
    names1 = [f"s1_{i}" for i in range(1, 9)]
    if not cell("s1_1"):
        names1 = [f"s1r1_{i}" for i in range(1, 5)] + [f"s1r2_{i}" for i in range(5, 9)]
    grid_pics(doc, names1, width=2.6, cols=4)

    S(doc, "二、听对话，根据听到的顺序，用字母 A—E 给下列图片排序。")
    names2 = [f"s2_{x}" for x in "ABCDE"]
    if not cell("s2_A"):
        names2 = [f"s2_{i}" for i in range(1, 6)]
    grid_pics(doc, names2, labels=list("ABCDE"), width=2.8, cols=5)

    S(doc, "三、听句子，选择正确的答语，将序号写在题前括号内。")
    reply3 = [
        [("A", "Sure."), ("B", "Please be polite.")],
        [("A", "I don't like it."), ("B", "I am happy.")],
        [("A", "I can go to the park."), ("B", "My toy car is broken.")],
        [("A", "He is tall with short hair."), ("B", "He is angry.")],
        [("A", "Of course. Let's do it together."), ("B", "You can remember the first word.")],
    ]
    for i, opts in enumerate(reply3, 1):
        P(doc, f"(  ) {i}.", size=11, bold=True, after=1)
        for lab, t in opts:
            P(doc, f"      {lab}. {t}", size=10, after=1)

    footer(doc, 1)
    doc.add_page_break()

    S(doc, "四、听对话，判断句子正误，正确的写“T”，错误的写“F”。")
    for i, s in enumerate(
        [
            "Linda is happy today.",
            "Linda doesn't remember a poem.",
            "Mike needs to use Linda's pencil.",
            "They draw a picture to understand the maths problem.",
            "Mike is a polite boy.",
        ],
        1,
    ):
        P(doc, f"(  ) {i}. {s}", size=11, after=4)

    S(doc, "五、听短文，选择正确答案，将序号写在题前括号内。")
    q5 = [
        ("Where are the children?", [("A", "At home."), ("B", "In a park."), ("C", "At school.")]),
        (
            "How do the children feel?",
            [("A", "They are excited."), ("B", "They are happy."), ("C", "They are sad.")],
        ),
        (
            "Who is Miss Green?",
            [("A", "An English teacher."), ("B", "A Chinese teacher."), ("C", "A music teacher.")],
        ),
        (
            "Why is Amy surprised?",
            [("A", "She gets a gift."), ("B", "She gets 100 in exam."), ("C", "She gets a cake.")],
        ),
        (
            "Where is Mike's English book?",
            [("A", "It's in his desk."), ("B", "It's under the chair."), ("C", "It's in his bag.")],
        ),
    ]
    for i, (q, opts) in enumerate(q5, 1):
        P(doc, f"(  ) {i}. {q}", size=11, bold=True, after=2)
        for lab, t in opts:
            P(doc, f"      {lab}. {t}", size=10, after=1)

    P(doc, "笔试部分", size=13, bold=True, center=True, before=12, after=8)

    S(doc, "六、读单词，选出不同类的一项，将序号写在题前括号内。")
    odd = [
        "A. worried   B. happy   C. feeling   D. angry",
        "A. story   B. English   C. Chinese   D. maths",
        "A. stand   B. sweep   C. water   D. open",
        "A. polite   B. help   C. draw   D. remember",
        "A. try   B. wait   C. door   D. use",
    ]
    for i, line in enumerate(odd, 1):
        P(doc, f"(  ) {i}. {line}", size=11, after=4)

    S(doc, "七、看图读短语，选择与图片相符的短语，将序号写在括号内。")
    for line in [
        "A. talk to a friend",
        "B. a bad cold",
        "C. listen to music",
        "D. tidy up the table",
        "E. open the window",
    ]:
        P(doc, line, size=10, after=2)
    grid_pics(doc, [f"s7_{i}" for i in range(1, 6)], width=2.8, cols=5)

    footer(doc, 2)
    doc.add_page_break()

    S(doc, "八、读句子，选择正确的答语，将序号写在题前括号内。")
    for line in [
        "A. You should have a good rest.",
        "B. We play football together.",
        "C. Sure. Here you are.",
        "D. It's OK. I can fix it.",
        "E. I can't find my Chinese book.",
    ]:
        P(doc, line, size=10, after=1)
    qs8 = [
        "You look worried. What's the matter?",
        "I have a stomachache.",
        "May I have a cupcake?",
        "What do you do with your friend?",
        "I'm sorry I broke your toy.",
    ]
    for i, q in enumerate(qs8, 1):
        P(doc, f"(  ) {i}. {q}", size=11, after=4)

    S(doc, "九、Max 一家在参观博物馆，他们在说些什么？读对话，从方框中选择合适的句子补全对话，将序号写在横线上。")
    for line in [
        "A. May I use your notebook?",
        "B. What's the matter?",
        "C. What is this?",
        "D. You can draw it down on your notebook.",
        "E. Thanks a lot!",
    ]:
        P(doc, line, size=10, after=1)
    for line in [
        "Max: Dad, 1. ________",
        "Dad: It's a rock (岩石).",
        "Max: It looks beautiful! I want to remember it.",
        "Dad: 2. ________",
        "Max: Oh, no!",
        "Mum: 3. ________",
        "Max: I can't find my notebook.",
        "Sister: Don't worry. I have my notebook.",
        "Max: 4. ________",
        "Sister: Sure. Here you are.",
        "Max: 5. ________ I can remember this rock when I go back home!",
    ]:
        P(doc, line, size=11, after=3)

    S(doc, "十、看图读故事，根据故事内容，用数字 1—6 给下列图片排序。")
    passage = (
        "It was sports day. Wilma was in the long jump. She came first. She was happy. "
        "Wilf and Chip were in a race. They came second. They were excited. "
        "Dad was in the egg and spoon race. He came third. Dad was surprised. "
        "Mum was in a race. But Mum came last. She was unhappy. "
        "Kids have a surprise for Mum. Mum is happy now."
    )
    P(doc, passage, size=11, after=8, indent=0.74)
    grid_pics(doc, [f"s10_{i}" for i in range(1, 7)], width=2.4, cols=3)

    footer(doc, 3)
    doc.add_page_break()

    S(doc, "十一、以下是 Happy Park 的游园须知。阅读须知内容，完成任务。")
    P(doc, "Welcome to Happy Park", size=12, bold=True, center=True, after=4)
    P(doc, "Open Time: 8:30–18:30 (Tue.–Sun.)", size=11, after=2)
    P(doc, "Ticket Price: Adult (成人) ¥20    Kid ¥10", size=11, after=6)
    park = cell("park_icons")
    if park:
        add_pic(doc, park, width_cm=14)
    P(doc, "You can:", size=11, bold=True, after=2)
    P(doc, "· fly kites    · draw pictures    · have a picnic    · play games", size=10, after=4)
    P(doc, "You can't:", size=11, bold=True, after=2)
    P(doc, "· climb the trees    · swim in the lake    · play football on the grass    · ride bikes", size=10, after=4)
    P(doc, "Tips: We have a snack bar. You can buy food and drinks.", size=10, after=8)

    P(doc, "任务一：根据须知内容，选择正确答案，将序号写在题前括号内。", size=11, bold=True, after=6)
    park_qs = [
        ("When does the park open?", [("A", "On Saturday."), ("B", "On Monday."), ("C", "At 7:00.")]),
        (
            "Mike is 9 years old. He wants to go to Happy Park with his father. They need to pay (支付)",
            [("A", "20 yuan"), ("B", "10 yuan"), ("C", "30 yuan")],
        ),
        (
            "What can we do in the park?",
            [("A", "We can swim."), ("B", "We can draw flowers."), ("C", "We can ride a bike.")],
        ),
        (
            "What can we buy in the park?",
            [("A", "Kites."), ("B", "Some pictures."), ("C", "Some water.")],
        ),
        (
            "Which is right?",
            [
                ("A", "We can go to the park at weekends."),
                ("B", "We can't play games."),
                ("C", "We can play football on the grass."),
            ],
        ),
    ]
    for i, (q, opts) in enumerate(park_qs, 1):
        P(doc, f"(  ) {i}. {q}", size=11, bold=True, after=2)
        for lab, t in opts:
            P(doc, f"      {lab}. {t}", size=10, after=1)

    P(doc, "任务二：如果你在 Happy Park 看到了这样的事情，你会说些什么？请你写一写。", size=11, bold=True, before=10, after=6)
    for name in ["w_1", "w_2", "w_3", "write"]:
        p = cell(name)
        if p:
            add_pic(doc, p, width_cm=5.5)
    P(doc, "1. _________________________________________________________________", size=11, after=4)
    P(doc, "2. _________________________________________________________________", size=11, after=4)
    P(doc, "3. _________________________________________________________________", size=11, after=8)

    footer(doc, 4)

    path = OUT / "U4_第四单元练习.docx"
    doc.save(path)
    shutil.copy(path, OUT / "U4_Unit4_Practice.docx")
    return path


def build_answers() -> Path:
    doc = new_doc()
    P(doc, "英语 · 四上 · 第四单元练习 · 参考答案", size=16, bold=True, center=True, after=6)
    P(doc, "（听力答案据《四上听力材料》脚本核对；图题答案对照原卷附图）", size=10, center=True, after=12)

    S(doc, "听力部分")
    P(doc, "一、听句子判断 T/F（材料脚本）：", size=11, bold=True, after=4)
    P(
        doc,
        "听力原文：1. Sometimes I feel excited.  2. My cat is white with a long tail.  "
        "3. The girl is ill today.  4. I often read books with my friends.  "
        "5. We should stand in line.  6. Could I have an apple, please?  "
        "7. May I use your eraser?  8. You should have a good rest.",
        size=9,
        after=4,
    )
    P(doc, "答案：以课堂录音 + 卷面图片是否相符为准（打印卷附图同原卷）。常见卷面参考：请家长按图核验。", size=10, after=6)

    P(doc, "二、听对话排序 A—E：", size=11, bold=True, after=4)
    P(
        doc,
        "材料顺序：A 找英语书 → B 在读故事 → C 大声读 → D 电话找 Mike（在睡觉）→ E 借铅笔。"
        "卷面排序：把听到的第 1–5 段依次标到对应图上（答案依图序填写，对照附图）。",
        size=10,
        after=6,
    )

    P(doc, "三、听选答语：1A  2B  3B  4A  5A", size=11, after=4)
    P(
        doc,
        "依据：1 open the window→Sure.  2 How do you feel→I am happy.  "
        "3 What's the matter / angry→My toy car is broken.  "
        "4 What does he look like→tall with short hair.  "
        "5 Can I fly a kite→Of course. Let's do it together.",
        size=9,
        after=6,
    )

    P(doc, "四、听对话判断：1F  2F  3F  4T  5T", size=11, after=4)
    P(
        doc,
        "依据：Linda: Not very happy.（非 happy）→1F；难的是 maths problem 不是 poem→2F；"
        "Mike 说 May I use your pen?（不是 pencil）→3F；续说 draw a picture 懂了→4T；礼貌请求→5T。",
        size=9,
        after=6,
    )

    P(doc, "五、听短文选择：1C  2B  3A  4A  5C", size=11, after=4)
    P(
        doc,
        "依据：back at school；very happy；Miss Green is their new English teacher；"
        "friends give her a gift；book is in Mike's schoolbag.",
        size=9,
        after=8,
    )

    S(doc, "笔试部分")
    P(doc, "六、异类词：1C  2A  3C  4A  5C", size=11, after=4)
    P(
        doc,
        "1 feeling 非心情形容词；2 story 非学科；3 water 非动作动词；4 polite 非动词；5 door 非动词。",
        size=9,
        after=6,
    )

    P(doc, "七、看图选短语：对照卷面附图勾选（选项 A–E）。", size=11, after=2)
    P(doc, "建议家长按图核对：对话/感冒/听音乐/整理桌子/开窗。", size=9, after=6)

    P(doc, "八、选答语：1E  2A  3C  4B  5D", size=11, after=6)

    P(doc, "九、补全对话：1C  2D  3B  4A  5E", size=11, after=6)

    P(doc, "十、故事排序：按短文情节 1→6（Wilma 第一 → Wilf&Chip 第二 → Dad 第三 → Mum 最后 → 惊喜 → Mum 开心）。", size=11, after=6)

    P(doc, "十一·任务一：1A  2C  3B  4C  5A", size=11, after=4)
    P(
        doc,
        "1 周二至周日含周六；2 小孩10+成人20=30；3 可画画不可游泳/骑车；"
        "4 小吃部可买水；5 周末可去；不可踢球；可玩游戏。",
        size=9,
        after=6,
    )
    P(doc, "十一·任务二：开放作答。例：Don't climb the trees. / Please don't swim in the lake. / You can fly a kite here.", size=10, after=8)

    path = OUT / "U4_第四单元练习_答案.docx"
    doc.save(path)
    shutil.copy(path, OUT / "U4_Unit4_Practice_Answers.docx")
    return path


def export_pdfs(paths: list[Path]) -> None:
    import win32com.client

    word = win32com.client.Dispatch("Word.Application")
    word.Visible = False
    try:
        for docx in paths:
            pdf = docx.with_suffix(".pdf")
            d = word.Documents.Open(str(docx))
            d.SaveAs(str(pdf), FileFormat=17)
            d.Close(False)
            print("pdf", pdf.name)
    finally:
        word.Quit()


def build_scan_pdf() -> Path | None:
    """Optional: multi-page PDF from hq page photos for pixel对照打印."""
    try:
        from PIL import Image
    except ImportError:
        return None
    pages = []
    for n in range(1, 5):
        p = HQ / f"hq_p{n}.jpg"
        if p.exists():
            pages.append(Image.open(p).convert("RGB"))
    if not pages:
        return None
    out = OUT / "U4_Unit4_Practice_Scan.pdf"
    pages[0].save(out, save_all=True, append_images=pages[1:], resolution=150)
    shutil.copy(out, OUT / "U4_第四单元练习_原卷扫描.pdf")
    print("scan pdf", out.name)
    return out


if __name__ == "__main__":
    paper = build_paper()
    ans = build_answers()
    print("docx", paper.name, ans.name)
    build_scan_pdf()
    export_pdfs(
        [
            paper,
            OUT / "U4_Unit4_Practice.docx",
            ans,
            OUT / "U4_Unit4_Practice_Answers.docx",
        ]
    )
    # Chinese-named pdfs
    shutil.copy(OUT / "U4_Unit4_Practice.pdf", OUT / "U4_第四单元练习.pdf")
    shutil.copy(OUT / "U4_Unit4_Practice_Answers.pdf", OUT / "U4_第四单元练习_答案.pdf")
    print("done")
